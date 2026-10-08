import json
import zipfile
import argparse
import os
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
MODS = Path(os.environ["ENTITYBATTLE_DOC_MODS"]) if os.environ.get("ENTITYBATTLE_DOC_MODS") else None
OUTPUT = ROOT / "EntityBattle原版与整合包生物属性表.docx"
DATA = ROOT / "src/main/resources/data/entitybattle"


def jar(pattern):
    if MODS is None:
        raise ValueError("Supply --mods or ENTITYBATTLE_DOC_MODS for Chinese source translations")
    matches = list(MODS.glob(pattern))
    if not matches:
        raise FileNotFoundError(pattern)
    return matches[0]


def json_in_zip(path, member):
    with zipfile.ZipFile(path) as archive:
        return json.loads(archive.read(member).decode("utf-8-sig"))


def set_font(run, size=8, bold=False, color="000000"):
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def cell_text(cell, first, second=None, centered=False):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    set_font(paragraph.add_run(str(first)), 8.5)
    if second:
        run = paragraph.add_run("\n" + str(second))
        set_font(run, 7.5, color="586975")


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), color)
    tc_pr.append(node)


def cell_margins(cell, vertical=90):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, amount in (("top", vertical), ("bottom", vertical), ("start", 90), ("end", 90)):
        node = OxmlElement("w:" + side)
        node.set(qn("w:w"), str(amount))
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tc_pr.append(mar)


def table_borders(table):
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement("w:" + side)
        for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            node.set(qn("w:" + key), value)
        borders.append(node)
    table._tbl.tblPr.append(borders)


def load_species():
    species = {}
    with zipfile.ZipFile(jar("*Cobblemon-neoforge*.jar")) as archive:
        for member in archive.namelist():
            if member.startswith("data/cobblemon/species/") and member.endswith(".json"):
                name = Path(member).stem
                species["cobblemon:" + name] = json.loads(archive.read(member).decode("utf-8-sig"))
    for path in (DATA / "species").glob("*.json"):
        species["entitybattle:" + path.stem] = json.loads(path.read_text(encoding="utf-8"))
    return species


def load_locales():
    locale = {}
    sources = [
        ("twilightforest", "*twilightforest*.jar"),
        ("aether", "*aether-1.21.1-1.5.10*.jar"),
        ("deep_aether", "*deep_aether*.jar"),
        ("kaleidoscope_twilight", "*kaleidoscope_twilight*.jar"),
        ("artifacts", "*artifacts-neoforge*.jar"),
        ("legendarymonuments", "*legendarymonuments*.jar"),
        ("cobblemon", "*Cobblemon-neoforge*.jar"),
    ]
    for namespace, pattern in sources:
        try:
            locale.update(json_in_zip(jar(pattern), f"assets/{namespace}/lang/zh_cn.json"))
        except (FileNotFoundError, KeyError):
            pass
    locale.update(json.loads((ROOT / "src/main/resources/assets/entitybattle/lang/zh_cn.json").read_text(encoding="utf-8")))
    return locale


def data_rows():
    locale = load_locales()
    species = load_species()
    paths = sorted((DATA / "battle_profiles").glob("*.json"))
    profiles = []
    for path in paths:
        value = json.loads(path.read_text(encoding="utf-8"))
        profiles.extend(value if isinstance(value, list) else [value])
    groups = {"twilightforest": "暮色森林", "aether": "天境", "deep_aether": "深入天境",
              "artifacts": "其他内容模组", "legendarymonuments": "其他内容模组",
              "minecraft": "原版生物", "touhou_little_maid": "车万女仆野外生物",
              "kaleidoscope_twilight": "森罗物语暮色"}
    rows = []
    for p in profiles:
        ns, entity = p["entity"].split(":", 1)
        sn, name = p["species"].split(":", 1)
        spec = species[p["species"]]
        pokemon_cn = locale.get(f"{sn}.species.{''.join(char for char in spec['name'].lower() if char.isalnum())}.name",
                                {"cow": "牛（独立物种）", "zombie": "僵尸（独立物种）"}.get(name, name))
        source_cn = locale.get(f"entity.{ns}.{entity}", pokemon_cn)
        types = "/".join(locale.get("cobblemon.type." + t, t) for t in
                         (spec.get("primaryType"), spec.get("secondaryType")) if t)
        abilities = "、".join(locale.get("cobblemon.ability." + ability.removeprefix("h:"), ability)
                             for ability in spec.get("abilities", []))
        learning_moves = "、".join(
            f'{entry.split(":", 1)[0]}级 {locale.get("cobblemon.move." + entry.split(":", 1)[1], entry.split(":", 1)[1])}'
            for entry in spec.get("moves", []) if entry.split(":", 1)[0].isdigit()
        )
        stats = "/".join(str(spec["baseStats"][stat]) for stat in
                         ("hp", "attack", "defence", "special_attack", "special_defence", "speed"))
        fixed_moves = "、".join(locale.get("cobblemon.move." + move, move) for move in p.get("moves", []))
        moves = ("野外固定四招 " + fixed_moves + "；学习表 " + learning_moves) if fixed_moves else learning_moves
        mode = "剧情后宝可梦" if p.get("boss") and ns in ("aether", "deep_aether") else (
            "自动宝可梦" if p.get("boss") else "直接宝可梦" if p["worldMode"] == "pokemon_entity" else "原生人工智能")
        if entity == "quest_ram":
            mode += "；任务领奖后可转换捕捉"
        rows.append({
            "group": groups[ns], "source": source_cn, "entity_id": p["entity"],
            "pokemon": pokemon_cn, "species_id": p["species"],
            "level": f'{p["level"]["min"]}-{p["level"]["max"]}',
            "catchable": "可捕捉" if p["catchable"] else "不可捕捉",
            "mode": mode, "boss": p.get("boss", False), "profile": p,
            "behavior": "被动" if p["worldBehavior"] == "passive" else "敌对",
            "types": types, "abilities": abilities, "stats": stats,
            "catch_rate": spec.get("catchRate", "-"),
            "exp": spec.get("baseExperienceYield", "-"),
            "moves": moves,
            "growth": {"medium_fast":"中速", "slow":"慢速"}[spec["experienceGroup"]],
            "ev": "、".join(f"{dict(hp='生命',attack='攻击',defence='防御',special_attack='特攻',special_defence='特防',speed='速度')[key]}+{value}" for key, value in spec.get("evYield", {}).items() if value) or "无",
            "height": spec.get("height", "-"),
            "weight": spec.get("weight", "-"),
            "number": spec.get("nationalPokedexNumber", "-"),
            "gender": "无性别" if spec.get("maleRatio") == -1.0 else
                      f'{spec.get("maleRatio", 0.5) * 100:.0f}%',
            "hitbox": spec.get("hitbox", {}),
        })
    return rows


def add_group_table(doc, rows):
    headers = ["来源生物", "宝可梦名称", "等级与捕捉", "属性与特性",
               "种族值\n生命/攻/防/特攻/特防/速", "捕获与成长"]
    widths = [4.9, 4.9, 2.9, 5.3, 5.4, 3.8]
    table = doc.add_table(rows=1, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table_borders(table)
    grid_widths = [int(width * 567) for width in widths]
    for column, width in zip(table._tbl.tblGrid.gridCol_lst, grid_widths):
        column.set(qn("w:w"), str(width))
    tbl_pr = table._tbl.tblPr
    tbl_width = tbl_pr.first_child_found_in("w:tblW")
    if tbl_width is None:
        tbl_width = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_width)
    tbl_width.set(qn("w:w"), str(sum(grid_widths)))
    tbl_width.set(qn("w:type"), "dxa")
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    header = table.rows[0]
    header._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for idx, title in enumerate(headers):
        cell_text(header.cells[idx], title, centered=True)
        header.cells[idx].width = Cm(widths[idx])
        header.cells[idx].paragraphs[0].paragraph_format.keep_with_next = True
        for run in header.cells[idx].paragraphs[0].runs:
            set_font(run, 7.7, bold=True, color="FFFFFF")
        shade(header.cells[idx], "27566F")
        cell_margins(header.cells[idx])
    for i, row in enumerate(rows):
        current_row = table.add_row()
        current_row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        cells = current_row.cells
        values = [
            (row["source"], "首领" if row["boss"] else "普通生物"),
            (row["pokemon"], "独立物种"),
            (row["level"], row["catchable"] + " · " + row["behavior"] + "\n" + row["mode"]),
            (row["types"], row["abilities"] + "\n招式 " + row["moves"]),
            (row["stats"], None),
            (str(row["catch_rate"]), "基础经验 " + str(row["exp"]) +
             "\n成长组 " + str(row["growth"]) + " · 努力值 " + str(row["ev"]) +
             "\n图鉴 " + str(row["number"]) + " · 雄性比例 " + str(row["gender"]) +
             "\n身高 " + str(row["height"]) + " · 体重 " + str(row["weight"]) +
             "\n碰撞箱 " + str(row["hitbox"].get("width", "-")) + " × " +
             str(row["hitbox"].get("height", "-"))),
        ]
        for idx, (main, sub) in enumerate(values):
            cell_text(cells[idx], main, sub, centered=idx in (2, 4, 5))
            cell_margins(cells[idx])
            cells[idx].width = Cm(widths[idx])
            if i % 2:
                shade(cells[idx], "F3F7F9")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def build(output=OUTPUT, pack_only=False):
    rows = data_rows()
    if pack_only:
        rows = [row for row in rows if not row["entity_id"].startswith("minecraft:")]
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.left_margin = section.right_margin = Cm(1)
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.1)
    for name in ("Normal", "Title", "Heading 1"):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Normal"].font.size = Pt(9.5)
    doc.styles["Normal"].paragraph_format.space_after = Pt(5)
    doc.styles["Title"].font.size = Pt(21)
    doc.styles["Title"].font.bold = True
    title_ppr = doc.styles["Title"]._element.get_or_add_pPr()
    for border in title_ppr.findall(qn("w:pBdr")):
        title_ppr.remove(border)
    doc.styles["Heading 1"].font.size = Pt(13)
    doc.styles["Heading 1"].font.bold = True

    title = "整合包生物宝可梦属性表" if pack_only else "原版与整合包生物宝可梦属性表"
    doc.add_paragraph(title, style="Title")
    p = doc.add_paragraph()
    set_font(p.add_run("生物宝可梦接入模组 · 我的世界 1.21.1 · 方可梦 1.8.1 · 更新 2026-10-08"),
             9.5, color="586975")
    doc.add_paragraph(
        f"本表从源码读取当前 {len(rows)} 种生物配置；完整清单共 169 种，正式首领 15 种。"
        "每种生物都有独立物种；来源中文名来自整合包的中文本地化。"
        "具有主人、背包和工作状态的女仆本体未接入。")
    doc.add_paragraph(
        "普通生物保留原生实体、人工智能和交互。正式首领默认满足原作开战条件后自动成为宝可梦；"
        "没有前置条件的首领（含向日葵）生成后即转换。天境与深入天境保留对话、勋章和唤醒要求。"
        "默认关闭普通生物快捷临时挑战，转换器仍可使用。"
        "表中种族值是基础值，个体值、性格与等级会影响实际面板数值。"
        "不同来源生物不共用物种标识。")

    doc.add_heading("两种接入方式", level=1)
    doc.add_paragraph(
        "原生生物模式：来源生物在大世界运行自己的人工智能、攻击、交互及动画。"
        "默认用转换器右击目标，将其永久变为真正的野生宝可梦实体。"
        "若开启快捷临时挑战，可对准目标临时转换对战：未捕捉时战后还原原生生物，"
        "捕捉后进入队伍或电脑，战败时尝试通过原生死亡流程结算。"
        "谜题羊不是首领，保持原生生物与原生人工智能；完成十六色羊毛任务并实际领奖后，玩家可选择转换捕捉。")
    doc.add_paragraph(
        "快捷挑战默认关闭。首领世界模式默认为宝可梦实体；切换为原生生物模式可保留原生首领。"
        "所有转换入口共用剧情资格检查。本批复杂首领保存原生完整状态，完整胜利执行原生死亡回调；"
        "中断不发钥匙、不解锁、不发首领奖励。临时转换中断恢复来源，自动模式中断保持宝可梦。")
    doc.add_paragraph(
        "直接宝可梦模式：生物生成时即替换成真正的宝可梦实体，"
        "默认使用方可梦行为与完整对战、捕捉、队伍流程。原生人工智能、特殊交互和死亡事件不自动保留，"
        "需要来源行为时必须注册对应适配器。两种模式都复用来源生物的模型与基础渲染，"
        "物种数据使用各自独立的设定文件。捕捉后的谜题羊使用方可梦行为，保留羊毛颜色，不再运行原生任务人工智能。")

    doc.add_heading("如何修改及复用", level=1)
    doc.add_paragraph(
        "每一行都有对应的独立物种文件。编辑其六项种族值、主副属性、特性、学习表、捕获率、"
        "基础经验和成长组，只影响这一种生物。"
        "本包已提供逐种写好的初始配置，接入作者可复制并修改，不必重新创建物种。")
    doc.add_paragraph(
        "接入行为由接入配置决定：可调整等级范围、捕捉资格、世界模式、来源掉落、原版经验、"
        "战斗道具与随从阵容。技术字段和具体路径请查看项目中的中文开发接入文档。"
        "修改后重新构建模组，也可用数据包覆盖；已有个体保存的招式、等级和道具不保证自动更新。")
    doc.add_paragraph(
        "不可捕捉的包括正式首领，以及原版远古守卫者、铁傀儡、劫掠兽、村民、流浪商人、"
        "监守者与凋灵。末影龙的多部件与死亡流程需要专门适配，当前未接入；巨人和幻术师属于未自然生成的特殊实体。"
        "谜题羊可捕捉，必须先完成原生任务领奖。妖精的颜色、幼体外观和点数掉落，"
        "以及复杂首领阶段、任务回调和掉落仍需实机验证。")

    doc.add_heading("转化粉与生存获取", level=1)
    doc.add_paragraph("作者为 RunWater1。普通僵尸现可捕捉，捕获率为180。转换器可用两份紫水晶碎片、两份红石和一份铁锭合成；整合包可按配方标识覆盖。捕捉并培养来源宝可梦，放出后用一份暮色转化粉右击，再在原生队伍详情确认进化。下表是本模组新增路线。")
    doc.add_paragraph("原生进化保留等级、个体值、努力值、性格、昵称、主人、精灵球、携带物与已学招式，使用结果物种的属性、特性与模型。不附赠随从装备，不继承旧世界死亡身份。正式首领仍需整场胜利；谜题羊仍须羊毛任务领奖。召唤类生物有的已有原作获取途径，本表补充稳定培养路线。详细修改方法见项目中文生存获取说明。")
    routes = json.loads((ROOT / "validation/creature-acquisition.json").read_text(encoding="utf-8"))
    route_table = doc.add_table(rows=1, cols=4)
    route_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    route_table.autofit = False
    table_borders(route_table)
    route_widths = [4.3, 4.3, 2.4, 16.2]
    for idx, heading in enumerate(("来源宝可梦", "进化结果", "最低等级", "补充原因")):
        cell_text(route_table.rows[0].cells[idx], heading)
        cell_margins(route_table.rows[0].cells[idx], vertical=50)
        shade(route_table.rows[0].cells[idx], "27566F")
        for run in route_table.rows[0].cells[idx].paragraphs[0].runs:
            set_font(run, 8.5, True, "FFFFFF")
    route_table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for index, route in enumerate(routes):
        cells = route_table.add_row().cells
        cells[0]._tc.getparent().get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for idx, value in enumerate((route["sourceName"], route["targetName"], route["level"], route["reason"])):
            cell_text(cells[idx], value)
            cell_margins(cells[idx], vertical=50)
            cells[idx].width = Cm(route_widths[idx])
            if index % 2:
                shade(cells[idx], "F3F7F9")
    for idx, column in enumerate(route_table.columns):
        column.width = Cm(route_widths[idx])
        route_table.rows[0].cells[idx].width = Cm(route_widths[idx])

    order = ("原版生物", "暮色森林", "森罗物语暮色", "天境", "深入天境", "其他内容模组", "车万女仆野外生物")
    for group in order:
        group_rows = [r for r in rows if r["group"] == group]
        if not group_rows:
            continue
        doc.add_heading(f"{group}  {len(group_rows)} 种", level=1)
        add_group_table(doc, group_rows)

    doc.add_heading("首领与随从战斗装备", level=1)
    doc.add_paragraph("正式首领的野外固定四招见上表；随从的本场装备与招式覆盖见下表。战斗携带物自动追加掉落已关闭，来源原作战利品仍保留。幻影骑士使用原有六骑士，不另造五个随从。暮色领主后续制作已取消，训练家由新整合包其他模组负责。")
    locale = load_locales()
    by_source = {row["entity_id"]: row for row in data_rows()}
    def item_name(value):
        ns, key = value.split(":", 1)
        return locale.get(f"item.{ns}.{key}") or locale[f"block.{ns}.{key}"]
    equipment = doc.add_table(rows=1, cols=4)
    equipment.alignment = WD_TABLE_ALIGNMENT.CENTER
    equipment.autofit = False
    table_borders(equipment)
    widths = [4.0, 5.0, 6.0, 12.2]
    for idx, heading in enumerate(("首领", "本体持有物", "本场随从", "随从装备与固定四招")):
        cell_text(equipment.rows[0].cells[idx], heading)
        cell_margins(equipment.rows[0].cells[idx])
        shade(equipment.rows[0].cells[idx], "27566F")
        for run in equipment.rows[0].cells[idx].paragraphs[0].runs:
            set_font(run, 8.5, True, "FFFFFF")
    equipment.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for index, row in enumerate(r for r in rows if r["boss"]):
        p = row["profile"]
        h = p.get("bossBattle")
        roster = "无额外随从"
        extra = "—"
        if h:
            m = by_source[h["minionEntity"]]
            off = h["levelOffset"]
            roster = f'{h["minionCount"]}名{m["source"]}；{max(1,p["level"]["min"]+off)}–{min(100,p["level"]["max"]+off)}级'
            extra = item_name(h.get("minionHeldItem", m["profile"].get("heldItem"))) + "；" + "、".join(locale["cobblemon.move." + move] for move in h.get("minionMoves", m["profile"].get("moves", [])))
            if p["entity"] == "twilightforest:knight_phantom":
                roster = "原有六骑士中的另五名；" + roster.split("；",1)[1]
        new_row = equipment.add_row()
        new_row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for idx, value in enumerate((row["source"], item_name(p["heldItem"]), roster, extra)):
            cell_text(new_row.cells[idx], value)
            cell_margins(new_row.cells[idx])
            new_row.cells[idx].width = Cm(widths[idx])
            if index % 2:
                shade(new_row.cells[idx], "F3F7F9")
    for idx, column in enumerate(equipment.columns):
        column.width = Cm(widths[idx])
        equipment.rows[0].cells[idx].width = Cm(widths[idx])

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(footer.add_run("生物宝可梦接入模组  ·  " + title), 7.5, color="586975")
    doc.core_properties.title = title
    doc.core_properties.subject = "接入生物设定"
    doc.core_properties.author = "RunWater1"
    doc.save(output)
    print(json.dumps({"output": str(output), "rows": len(rows), "bytes": output.stat().st_size}, ensure_ascii=False))


def main():
    global MODS
    parser=argparse.ArgumentParser()
    parser.add_argument("--mods",type=Path,default=MODS)
    args=parser.parse_args()
    MODS=args.mods
    build()
    build(ROOT / "EntityBattle整合包生物属性表.docx", pack_only=True)

if __name__ == "__main__":
    main()
